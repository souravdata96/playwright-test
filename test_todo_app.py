from pathlib import Path
import shutil

from docx import Document
from docx.shared import Inches
from playwright.sync_api import Page, expect


TODO_INPUT = "What needs to be done?"
TODO_TOGGLE = "Toggle Todo"

ACTIONS = [
    ("add", "feed the dog"),
    ("check", "feed the dog"),
    ("add", "water the plants"),
    ("check", "water the plants"),
    ("add", "buy chocolate"),
    ("submit", None),
    ("add", "today"),
    ("uncheck", "water the plants"),
    ("uncheck", "feed the dog"),
    ("check", "water the plants"),
    ("active", None),
    ("click", "feed the dog"),
    ("check", "feed the dog"),
]


def capture_step(
    page, report: Document, screenshot_dir: Path, step_number: int, action: str, text: str | None
) -> None:
    description = action if text is None else f"{action} - {text}"
    filename = f"{step_number:02d}_{action}_{text or 'none'}.png".replace(" ", "_")
    screenshot_path = screenshot_dir / filename
    page.screenshot(path=str(screenshot_path))
    report.add_heading(f"Step {step_number:02d}: {description}", level=2)
    report.add_picture(str(screenshot_path), width=Inches(6))
    print(f"Step {step_number:02d}: {description}")


def run(page: Page) -> None:
    screenshot_dir = Path("test-results/screenshots")
    report_path = Path("test-results/test_execution_report.docx")
    screenshot_dir.mkdir(parents=True, exist_ok=True)
    report = Document()
    report.add_heading("Playwright Test Execution Report", level=1)
    report.add_paragraph("Each section shows the action performed and the resulting page state.")
    page.goto("https://demo.playwright.dev/todomvc/#/")
    todo_input = page.get_by_role("textbox", name=TODO_INPUT)
    todos = page.locator(".todo-list > li")
    active_link = page.get_by_role("link", name="Active")

    for step_number, (action, text) in enumerate(ACTIONS, start=1):
        if action == "add":
            todo_input.fill(text)
            todo_input.press("Enter")
            expect(page.get_by_text(text)).to_be_visible()
        elif action == "submit":
            todo_input.press("Enter")
            expect(todo_input).to_have_value("")
        elif action in ("check", "uncheck"):
            toggle = todos.filter(has_text=text).get_by_label(TODO_TOGGLE)
            toggle.evaluate("element => element.click()")
            page.goto("https://demo.playwright.dev/todomvc/#/")
            toggle = todos.filter(has_text=text).get_by_label(TODO_TOGGLE)
            if action == "check":
                expect(toggle).to_be_checked()
            else:
                expect(toggle).not_to_be_checked()
        elif action == "active":
            active_link.click()
            expect(page).to_have_url("https://demo.playwright.dev/todomvc/#/active")
            expect(todos).to_have_count(3)
        elif action == "click":
            page.get_by_text(text).click()
            expect(page.get_by_text(text)).to_be_visible()

        capture_step(page, report, screenshot_dir, step_number, action, text)
        page.wait_for_timeout(300)

    report.save(report_path)
    shutil.rmtree(screenshot_dir)
    print(f"Word report: {report_path}")


def test_todo_app(page: Page):
    run(page)
